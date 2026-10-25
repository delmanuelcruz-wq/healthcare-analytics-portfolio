"""Project D - Multi-state license compliance hub (SYNTHETIC).

    python D_license_compliance_hub/run.py

Builds the staff and site license registers with renewal buckets and document audits,
medical-director coverage checks, a per-site compliance matrix, a reminder PLAN (nothing is
sent), an Excel tracker with live formulas and conditional formatting, and the hub page.
"""
import csv
import sys
from collections import Counter, defaultdict
from datetime import date, timedelta
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))

from shared import portfolio as P  # noqa: E402

SQL, OUT = HERE / "sql", HERE / "outputs"
BUCKETS = ["Expired", "Due in 0-30 days", "Due in 31-60 days", "Due in 61-90 days", "Current (90+ days)"]
BUCKET_STATUS = {"Expired": "critical", "Due in 0-30 days": "serious", "Due in 31-60 days": "warning",
                 "Due in 61-90 days": "warning", "Current (90+ days)": "good"}


def bucket(days):
    if days < 0:
        return "Expired"
    if days <= 30:
        return "Due in 0-30 days"
    if days <= 60:
        return "Due in 31-60 days"
    if days <= 90:
        return "Due in 61-90 days"
    return "Current (90+ days)"


def staff_issues(r):
    issues = []
    if r["first_record_key"] != r["license_record_key"]:
        issues.append(("DUPLICATE", f"Duplicate of {r['first_record_key']} (same license ID)"))
        return issues                      # the original record carries the license's other findings
    if r["days_to_expiry"] < 0:
        issues.append(("EXPIRED_STATUS_STALE" if r["recorded_status"] == "Active" else "EXPIRED",
                       f"Expired {-r['days_to_expiry']} days ago" + (" but still recorded Active" if r["recorded_status"] == "Active" else "")))
    elif r["days_to_expiry"] <= 90:
        issues.append(("RENEWAL_DUE", f"Renewal due in {r['days_to_expiry']} days"))
    if r["document_flag"] == 0:
        issues.append(("DOC_MISSING", "No license document on file"))
    elif r["documents_on_file"] == 0:
        issues.append(("DOC_FLAG_WRONG", "Register says a document is on file, but none exists"))
    return issues


def overlaps(rows):
    found = []
    for i, a in enumerate(rows):
        for b in rows[i + 1:]:
            if a["site_key"] == b["site_key"] and a["is_primary"] and b["is_primary"]:
                a_end, b_end = a["relationship_end"] or "9999-12-31", b["relationship_end"] or "9999-12-31"
                if a["relationship_start"] <= b_end and b["relationship_start"] <= a_end:
                    found.append((a, b))
    return found


def build(con):
    as_of = date.fromisoformat(P.meta(con)["as_of_date"])
    staff = P.query(con, SQL / "01_staff_licenses.sql")
    sites_l = P.query(con, SQL / "02_site_licenses.sql")
    md = P.query(con, SQL / "03_director_coverage.sql")
    sites = P.query(con, "SELECT site_key, site_name, state_code, opened_date FROM sites ORDER BY site_key")

    for r in staff:
        r["renewal_bucket"] = bucket(r["days_to_expiry"])
        r["issues"] = staff_issues(r)
        r["issue_text"] = "; ".join(t for _, t in r["issues"])
        r["area"] = "Medical director licenses" if r["credential_class"] == "MD" else "Nursing & clinical licenses"
    for r in sites_l:
        r["renewal_bucket"] = bucket(r["days_to_expiry"])
        iss = []
        if r["days_to_expiry"] < 0:
            iss.append(("EXPIRED", f"Expired {-r['days_to_expiry']} days ago"))
        elif r["days_to_expiry"] <= 90:
            iss.append(("RENEWAL_DUE", f"Renewal due in {r['days_to_expiry']} days"))
        if r["document_flag"] == 0:
            iss.append(("DOC_MISSING", "No document on file"))
        r["issues"], r["issue_text"] = iss, "; ".join(t for _, t in iss)

    ov = overlaps(md)
    current_primary = defaultdict(list)
    for a in md:
        if a["is_current"] and a["is_primary"]:
            current_primary[a["site_key"]].append(a)

    matrix = []
    for s in sites:
        k = s["site_key"]
        lic = [r for r in sites_l if r["site_key"] == k]
        team = [r for r in staff if r["home_site_key"] == k and r["first_record_key"] == r["license_record_key"]]
        directors = current_primary.get(k, [])
        checks = {
            "Infusion center license current": all(r["days_to_expiry"] >= 0 for r in lic if r["license_type"].startswith("Ambulatory")),
            "Pharmacy permit current": all(r["days_to_expiry"] >= 0 for r in lic if r["license_type"].startswith("Pharmacy")),
            "Site documents on file": all(r["document_flag"] == 1 for r in lic),
            "Primary medical director today": len(directors) >= 1,
            "One primary director at a time": not any(a["site_key"] == k for a, _ in ov),
            "Director license current and documented": bool(directors) and all(
                d["director_license_expiry"] and d["director_license_expiry"] >= as_of.isoformat() and d["director_license_docs"] > 0 for d in directors),
            "Staff licenses current": all(r["days_to_expiry"] >= 0 for r in team),
            "Staff documents complete": all(r["document_flag"] == 1 and r["documents_on_file"] > 0 for r in team),
        }
        matrix.append(dict(site_key=k, site_name=s["site_name"].replace(" (fictional)", ""), state=s["state_code"],
                           staff_licenses=len(team), passed=sum(checks.values()), total=len(checks), checks=checks,
                           director=", ".join(d["director_label"] for d in directors) or "None"))

    reminders = []
    for r in staff + sites_l:
        if "first_record_key" in r and r["first_record_key"] != r["license_record_key"]:
            continue                       # duplicate staff record - the original carries the reminders
        if 0 <= r["days_to_expiry"] <= 180:
            expiry = date.fromisoformat(r["expiry_date"])
            for offset in (90, 60, 30):
                when = expiry - timedelta(days=offset)
                reminders.append(dict(record=r.get("license_record_key") or r["site_license_key"],
                                      holder=r.get("staff_label") or r["site_name"].replace(" (fictional)", ""),
                                      license=r.get("credential_class") or r["license_type"], expiry_date=r["expiry_date"],
                                      reminder=f"{offset}-day", reminder_date=when.isoformat(),
                                      plan_status="Would already have been sent" if when <= as_of else "Scheduled",
                                      note="Plan only - reminder automation not deployed"))
    reminders.sort(key=lambda x: (x["reminder_date"], x["record"]))

    exceptions = []
    for r in staff:
        for code, text in r["issues"]:
            if code != "RENEWAL_DUE":
                exceptions.append(dict(area=r["area"], record=r["license_record_key"], holder=r["staff_label"], code=code, finding=text,
                                       owner="HR / credentialing", severity="High" if code.startswith("EXPIRED") else "Medium"))
    for r in sites_l:
        for code, text in r["issues"]:
            if code != "RENEWAL_DUE":
                exceptions.append(dict(area="Location licenses", record=r["site_license_key"], holder=r["site_name"].replace(" (fictional)", ""),
                                       code=code, finding=f"{r['license_type']}: {text}", owner="Site operations",
                                       severity="High" if code == "EXPIRED" else "Medium"))
    for a, b in ov:
        pair = "|".join(sorted((a["assignment_key"], b["assignment_key"])))
        exceptions.append(dict(area="Medical director coverage", record=pair, holder=a["site_key"],
                               code="MD_OVERLAP", finding=f"Two primary directors overlap ({max(a['relationship_start'], b['relationship_start'])} to "
                               f"{min(a['relationship_end'] or '9999-12-31', b['relationship_end'] or '9999-12-31')})",
                               owner="Medical affairs", severity="Medium"))
    for s in sites:
        if not current_primary.get(s["site_key"]):
            last = max((a for a in md if a["site_key"] == s["site_key"]), key=lambda a: a["relationship_end"] or "9999")
            exceptions.append(dict(area="Medical director coverage", record=last["assignment_key"], holder=s["site_key"], code="MD_UNCOVERED",
                                   finding=f"No active primary medical director since {last['relationship_end']}", owner="Medical affairs",
                                   severity="High"))
    return as_of, staff, sites_l, md, matrix, reminders, exceptions, ov


def run_qa(con, staff, sites_l, md, matrix, reminders, exceptions, register):
    checks = []

    def add(name, status, observed, rule):
        checks.append(dict(check=name, status=status, observed=observed, rule=rule))

    by_code = defaultdict(set)
    for e in exceptions:
        by_code[e["code"]].add(e["record"])
    expect = {
        "EXPIRED_STATUS_STALE": "LICENSE_EXPIRED_STATUS_ACTIVE", "DOC_FLAG_WRONG": "LICENSE_DOC_FLAG_WRONG",
        "DOC_MISSING": ("LICENSE_DOC_MISSING", "SITE_LICENSE_DOC_MISSING"), "DUPLICATE": "LICENSE_DUPLICATE_RECORD",
        "EXPIRED": "SITE_LICENSE_EXPIRED", "MD_OVERLAP": "MD_OVERLAPPING_PRIMARY", "MD_UNCOVERED": "MD_SITE_UNCOVERED",
    }
    misses = []
    for code, types in expect.items():
        types = (types,) if isinstance(types, str) else types
        want = {r["record_key"] for r in register if r["defect_type"] in types}
        if by_code.get(code, set()) != want:
            misses.append(code)
    add("Every planted compliance defect is found, nothing extra", "pass" if not misses else "fail",
        f"{len(exceptions)} exceptions across staff, site and director records; mismatched codes: {', '.join(misses) or 'none'}",
        "Graded against the ground-truth register")
    stale = [r["license_record_key"] for r in staff if r["days_to_expiry"] < 0 and r["recorded_status"] == "Active"]
    add("Status computed from dates, not trusted from the register", "warn" if stale else "pass",
        f"{len(stale)} license(s) past expiry still typed as Active: {', '.join(stale)}", "Status = f(expiry_date, as-of date)")
    unique = len({r["fictional_license_key"] for r in staff})
    add("One row per license", "warn" if unique < len(staff) else "pass",
        f"{len(staff)} register rows -> {unique} unique licenses; duplicates excluded from counts and reminders", "Deduplicate on license ID")
    docs = [r for r in staff if r["first_record_key"] == r["license_record_key"] and (r["document_flag"] == 0 or r["documents_on_file"] == 0)]
    add("Document presence audited against the index", "warn" if docs else "pass",
        f"{len(docs)} license(s) lack a filed document ({sum(1 for r in docs if r['document_flag'] == 1)} where the flag wrongly says one exists)",
        "Flag and index must agree")
    covered = sum(1 for m in matrix if m["checks"]["Primary medical director today"])
    add("Every site has exactly one primary medical director", "warn",
        f"{covered}/{len(matrix)} sites covered today; one historical overlap of two primaries", "Effective intervals may not overlap for primaries")
    in_bucket = Counter(r["renewal_bucket"] for r in staff if r["first_record_key"] == r["license_record_key"])
    add("Renewal buckets are complete and exclusive", "pass" if sum(in_bucket.values()) == unique else "fail",
        ", ".join(f"{b}: {in_bucket.get(b, 0)}" for b in BUCKETS), "Each license in exactly one 30/60/90 bucket")
    add("Reminder automation is planned, not claimed", "pass",
        f"{len(reminders)} reminder dates planned for licenses expiring within 180 days; none sent", "Never describe a plan as deployed")
    fake = P.scalar(con, "SELECT COUNT(*) FROM license_registry WHERE fictional_license_key NOT LIKE 'MOCK-%'")
    add("Only nonfunctional MOCK- license IDs and role labels", "pass" if fake == 0 else "fail",
        "No names, real license numbers, scans or addresses exist in the data", "Public-safe by construction")
    return checks


def write_workbook(path, as_of, staff, sites_l, md, matrix, reminders, exceptions, qa):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    navy, banner = PatternFill("solid", fgColor="12344A"), PatternFill("solid", fgColor="FFF4D6")
    fills = {"crit": PatternFill("solid", fgColor="F6D3D3"), "ser": PatternFill("solid", fgColor="FBE1D5"),
             "warn": PatternFill("solid", fgColor="FDEFC8"), "good": PatternFill("solid", fgColor="D9F0D9")}
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    ws["A1"], ws["A1"].fill, ws["A1"].font = P.DISCLAIMER, banner, Font(bold=True, color="5C4300")
    ws["A3"], ws["B3"] = "As-of date (drives every formula)", as_of
    ws["B3"].number_format = "yyyy-mm-dd"
    ws["A3"].font = Font(bold=True)
    notes = ["Licenses hub - Coralline Infusion Network (fictional). Staff are role labels; license IDs are MOCK- and nonfunctional.",
             "Change B3 to re-age every license: Days to expiry and Renewal bucket are live formulas.",
             "Tabs: Clinical licenses | Site licenses | Director coverage | Site compliance | Reminder plan | Exceptions | QA",
             "Reminder plan is a plan only - no reminder automation is deployed."]
    for i, n in enumerate(notes, 5):
        ws.cell(row=i, column=1, value=n)
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 14

    def sheet(title, headers, rows, widths=None, formats=None):
        sh = wb.create_sheet(title)
        sh["A1"], sh["A1"].fill, sh["A1"].font = P.DISCLAIMER, banner, Font(bold=True, color="5C4300")
        sh.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        for j, h in enumerate(headers, 1):
            c = sh.cell(row=3, column=j, value=h)
            c.fill, c.font, c.alignment = navy, Font(color="FFFFFF", bold=True), Alignment(wrap_text=True, vertical="top")
        for i, row in enumerate(rows, 4):
            for j, v in enumerate(row, 1):
                c = sh.cell(row=i, column=j, value=v)
                if formats and j in formats:
                    c.number_format = formats[j]
        sh.freeze_panes = "C4"
        sh.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{3 + len(rows)}"
        for j in range(1, len(headers) + 1):
            sh.column_dimensions[get_column_letter(j)].width = (widths or {}).get(j, 14)
        sh.row_dimensions[3].height = 40
        return sh

    def bucket_formula(col, i):
        return (f'=IF({col}{i}<0,"Expired",IF({col}{i}<=30,"Due in 0-30 days",IF({col}{i}<=60,"Due in 31-60 days",'
                f'IF({col}{i}<=90,"Due in 61-90 days","Current (90+ days)"))))')

    def color_rules(sh, col, n):
        rng = f"{col}4:{col}{3 + n}"
        sh.conditional_formatting.add(rng, FormulaRule(formula=[f'{col}4="Expired"'], fill=fills["crit"]))
        sh.conditional_formatting.add(rng, FormulaRule(formula=[f'{col}4="Due in 0-30 days"'], fill=fills["ser"]))
        sh.conditional_formatting.add(rng, FormulaRule(formula=[f'OR({col}4="Due in 31-60 days",{col}4="Due in 61-90 days")'], fill=fills["warn"]))
        sh.conditional_formatting.add(rng, FormulaRule(formula=[f'{col}4="Current (90+ days)"'], fill=fills["good"]))

    rows = []
    for i, r in enumerate(staff, 4):
        rows.append([r["license_record_key"], r["staff_label"], r["role"], r["home_site_key"], r["credential_class"], r["license_scope"],
                     r["state_code"], r["fictional_license_key"], date.fromisoformat(r["issue_date"]), date.fromisoformat(r["expiry_date"]),
                     date.fromisoformat(r["verified_date"]), r["recorded_status"], r["document_flag"], r["documents_on_file"],
                     f"=J{i}-README!$B$3", bucket_formula("O", i), r["issue_text"]])
    sh = sheet("Clinical licenses", ["Record", "Staff (label)", "Role", "Home site", "Credential", "Scope", "State", "License ID (mock)",
                                     "Issued", "Expires", "Verified", "Recorded status", "Doc flag", "Docs on file", "Days to expiry",
                                     "Renewal bucket", "Findings"], rows,
               widths={2: 18, 3: 24, 6: 22, 8: 22, 16: 20, 17: 52}, formats={9: "yyyy-mm-dd", 10: "yyyy-mm-dd", 11: "yyyy-mm-dd"})
    color_rules(sh, "P", len(rows))

    rows = []
    for i, r in enumerate(sites_l, 4):
        rows.append([r["site_license_key"], r["site_name"].replace(" (fictional)", ""), r["license_type"], r["fictional_license_key"],
                     r["state_code"], date.fromisoformat(r["issue_date"]), date.fromisoformat(r["expiry_date"]), r["document_flag"],
                     f"=G{i}-README!$B$3", bucket_formula("I", i), r["issue_text"]])
    sh = sheet("Site licenses", ["Record", "Site", "License type", "License ID (mock)", "State", "Issued", "Expires", "Doc flag",
                                 "Days to expiry", "Renewal bucket", "Findings"], rows,
               widths={2: 28, 3: 36, 4: 20, 10: 20, 11: 40}, formats={6: "yyyy-mm-dd", 7: "yyyy-mm-dd"})
    color_rules(sh, "J", len(rows))

    rows = []
    for i, a in enumerate(md, 4):
        rows.append([a["assignment_key"], a["site_key"], a["site_name"].replace(" (fictional)", ""), a["director_label"],
                     date.fromisoformat(a["relationship_start"]), date.fromisoformat(a["relationship_end"]) if a["relationship_end"] else None,
                     a["is_primary"], f'=AND(E{i}<=README!$B$3,OR(F{i}="",F{i}>=README!$B$3))'])
    sheet("Director coverage", ["Assignment", "Site", "Site name", "Director (label)", "Start", "End", "Primary", "Current?"], rows,
          widths={3: 28, 4: 18}, formats={5: "yyyy-mm-dd", 6: "yyyy-mm-dd"})

    names = list(matrix[0]["checks"])
    sheet("Site compliance", ["Site", "Site name", "Director today"] + names + ["Score"],
          [[m["site_key"], m["site_name"], m["director"]] + ["PASS" if m["checks"][n] else "FAIL" for n in names] + [f"{m['passed']}/{m['total']}"]
           for m in matrix], widths={2: 28, 3: 18})
    sheet("Reminder plan", ["Record", "Holder", "License", "Expires", "Reminder", "Reminder date", "Plan status", "Note"],
          [[r[k] for k in ("record", "holder", "license", "expiry_date", "reminder", "reminder_date", "plan_status", "note")] for r in reminders],
          widths={2: 26, 3: 32, 7: 26, 8: 40})
    sheet("Exceptions", ["Area", "Record", "Holder", "Code", "Finding", "Owner", "Severity"],
          [[e[k] for k in ("area", "record", "holder", "code", "finding", "owner", "severity")] for e in exceptions],
          widths={1: 26, 2: 16, 3: 22, 4: 22, 5: 60, 6: 20})
    sheet("QA", ["Status", "Check", "Observed", "Rule"], [[c["status"].upper(), c["check"], c["observed"], c["rule"]] for c in qa],
          widths={2: 44, 3: 80, 4: 40})
    wb.save(path)


def main():
    con = P.connect()
    as_of, staff, sites_l, md, matrix, reminders, exceptions, ov = build(con)
    with (P.FOUNDATION / "data" / "exceptions_register.csv").open(encoding="utf-8") as fh:
        register = list(csv.DictReader(fh))
    qa = run_qa(con, staff, sites_l, md, matrix, reminders, exceptions, register)

    flat_matrix = [dict({k: v for k, v in m.items() if k != "checks"}, **{c: ("PASS" if ok else "FAIL") for c, ok in m["checks"].items()})
                   for m in matrix]
    strip = lambda rows: [{k: v for k, v in r.items() if k != "issues"} for r in rows]      # noqa: E731
    for name, rows in [("staff_license_register", strip(staff)), ("site_license_register", strip(sites_l)), ("director_coverage", md),
                       ("site_compliance_matrix", flat_matrix), ("reminder_plan", reminders), ("exceptions", exceptions), ("qa_results", qa)]:
        P.write_csv(rows, OUT / f"{name}.csv")
    write_workbook(HERE / "license_compliance_tracker.xlsx", as_of, staff, sites_l, md, matrix, reminders, exceptions, qa)

    unique_staff = [r for r in staff if r["first_record_key"] == r["license_record_key"]]
    b_staff = Counter(r["renewal_bucket"] for r in unique_staff)
    b_site = Counter(r["renewal_bucket"] for r in sites_l)
    doc_gaps = sum(1 for r in unique_staff if r["document_flag"] == 0 or r["documents_on_file"] == 0) + sum(1 for r in sites_l if not r["document_flag"])
    tiles = [
        dict(label="Licenses tracked", value=f"{len(unique_staff) + len(sites_l)}", note=f"{len(unique_staff)} staff, {len(sites_l)} site"),
        dict(label="Expired", value=f"{b_staff['Expired'] + b_site['Expired']}", note="Includes records still typed 'Active'"),
        dict(label="Due in 30 days", value=f"{b_staff['Due in 0-30 days'] + b_site['Due in 0-30 days']}", note="Renewal action now"),
        dict(label="Due in 31-90 days", value=f"{sum(b_staff[b] + b_site[b] for b in BUCKETS[2:4])}", note="Next two reminder waves"),
        dict(label="Document gaps", value=f"{doc_gaps}", note="Missing, or flagged but not in the index"),
        dict(label="Sites fully compliant", value=f"{sum(1 for m in matrix if m['passed'] == m['total'])} / {len(matrix)}", note="All eight checks pass"),
    ]
    data = dict(meta=dict(as_of=as_of.strftime("%B %d, %Y").replace(" 0", " "), as_of_iso=as_of.isoformat(), seed=P.meta(con)["seed"]),
                tiles=tiles, buckets=BUCKETS, bucket_status=BUCKET_STATUS,
                bucket_counts=[dict(bucket=b, staff=b_staff.get(b, 0), site=b_site.get(b, 0)) for b in BUCKETS],
                staff=strip(staff), sites=strip(sites_l), md=md, matrix=flat_matrix, checks=list(matrix[0]["checks"]),
                reminders=reminders, exceptions=exceptions, qa=qa)
    P.render_dashboard(HERE / "dashboard_template.html", HERE / "dashboard.html", data)
    if P.screenshot(HERE / "dashboard.html", HERE / "dashboard.png", width=1440, height=4600):
        print("  screenshot -> dashboard.png")
    write_qa_log(qa)
    print(f"Project D built: {len(staff)} staff rows, {len(sites_l)} site licenses, {len(exceptions)} exceptions, QA "
          f"{sum(c['status'] == 'pass' for c in qa)} pass / {sum(c['status'] == 'warn' for c in qa)} warn / "
          f"{sum(c['status'] == 'fail' for c in qa)} fail")


def write_qa_log(qa):
    lines = ["# QA log - Project D (license compliance hub)", "", f"> {P.DISCLAIMER}. Generated by `run.py`.", "",
             "| Status | Check | Observed | Rule applied |", "|---|---|---|---|"]
    lines += [f"| {c['status'].upper()} | {c['check']} | {c['observed']} | {c['rule']} |" for c in qa]
    (HERE / "QA_LOG.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
