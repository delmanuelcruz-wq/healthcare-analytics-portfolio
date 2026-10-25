"""Tier 2 #16 tests: the seven-quarter trend counts distinct appointments and the workbook matches the SQL."""
import csv
import importlib.util
import sys
from pathlib import Path

import pytest

HERE = Path(__file__).resolve().parents[1]
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT))

from shared import portfolio as P   # noqa: E402

_spec = importlib.util.spec_from_file_location("tier2_p16_run", HERE / "run.py")
M = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(M)


@pytest.fixture(scope="module")
def con():
    c = P.connect()
    yield c
    c.close()


@pytest.fixture(scope="module")
def loaded(con):
    return M.load(con)


def test_seven_quarters_in_order(loaded):
    trend, _, _ = loaded
    assert [r["quarter"] for r in trend] == M.QUARTERS


def test_visits_never_exceed_line_items_and_reconcile(loaded):
    trend, _, _ = loaded
    for r in trend:
        assert r["treated_visits"] <= r["line_items"], r["quarter"]
        explained = r["line_items"] - r["waste_lines"] - r["given_lines_on_untreated_visits"] - r["extra_lines_same_visit"]
        assert explained == r["treated_visits"], r["quarter"]


def test_independent_python_matches_sql(con, loaded):
    trend, _, _ = loaded
    assert M.recompute_python(con) == {r["quarter"]: r["treated_visits"] for r in trend}


def test_planted_duplicate_line_adds_no_visit(con, loaded):
    trend, _, _ = loaded
    with (ROOT / "foundation" / "data" / "exceptions_register.csv").open(encoding="utf-8") as fh:
        dupes = [r["record_key"] for r in csv.DictReader(fh) if r["defect_type"] == "DUP_ADMIN_LINE"]
    mine = [k for k in dupes if P.scalar(con, "SELECT therapy_key FROM administrations WHERE admin_key = ?", (k,)) == M.THERAPY]
    assert mine, "expected at least one planted duplicate line for the chosen therapy"
    assert sum(r["exact_duplicate_lines"] for r in trend) == len(mine)
    for key in mine:   # removing the duplicate line must not change the visit count of its appointment
        appt = P.scalar(con, "SELECT appt_key FROM administrations WHERE admin_key = ?", (key,))
        assert P.scalar(con, """SELECT COUNT(*) FROM administrations WHERE appt_key = ? AND therapy_key = ?
                                AND administered_flag = 1 AND admin_key <> ?""", (appt, M.THERAPY, key)) >= 1


def test_workbook_matches_sql_and_uses_live_formulas(loaded):
    from openpyxl import load_workbook
    trend, _, _ = loaded
    wb = load_workbook(HERE / "cfo_seven_quarter_trend.xlsx")
    assert len(wb.sheetnames) == 1
    ws = wb.active
    assert "SYNTHETIC" in ws["A1"].value
    for i, r in enumerate(trend, 6):
        assert ws.cell(row=i, column=2).value == r["treated_visits"]
        assert ws.cell(row=i, column=6).value == r["line_items"]
    assert ws["C7"].value == "=B7/B6-1"
    assert ws["C14"].value == "=B12/B6-1"          # first-to-last
    assert ws._charts, "native Excel chart expected"
    footnote = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("Counts are")]
    assert footnote and "line-item" in footnote[0]


def test_qa_suite_has_no_failures(con, loaded):
    trend, bounds, _ = loaded
    assert not [c for c in M.run_qa(con, trend, bounds) if c["status"] == "fail"]
