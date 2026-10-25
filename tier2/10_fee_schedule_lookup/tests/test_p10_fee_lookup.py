"""Tier 2 #10 tests: one bucket per request, $0.00 is zero-pay, matches need all three keys, amounts labelled synthetic."""
import importlib.util
import sys
from collections import Counter
from pathlib import Path

import pytest

HERE = Path(__file__).resolve().parents[1]
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT))

from shared import portfolio as P   # noqa: E402

_spec = importlib.util.spec_from_file_location("tier2_p10_run", HERE / "run.py")
M = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(M)


@pytest.fixture(scope="module")
def con():
    c = P.connect()
    yield c
    c.close()


@pytest.fixture(scope="module")
def loaded(con):
    return M.load(con)


def test_every_request_has_exactly_one_known_bucket(con, loaded):
    tracker, _ = loaded
    assert len(tracker) == P.scalar(con, "SELECT COUNT(*) FROM lookup_requests")
    assert len({r["request_key"] for r in tracker}) == len(tracker)
    assert {r["bucket"] for r in tracker} <= set(M.BUCKETS)


def test_zero_pay_is_not_no_match(loaded):
    tracker, _ = loaded
    zero = [r for r in tracker if r["bucket"] == "Zero-pay ($0.00)"]
    assert zero and all(r["fictional_fee_amount_synthetic"] == 0 and r["matched_fee_row"] for r in zero)
    assert all(r["fictional_fee_amount_synthetic"] is None for r in tracker if r["bucket"] == "No match")


def test_match_requires_code_claim_type_and_date(con, loaded):
    tracker, _ = loaded
    for r in tracker:
        if r["bucket"] in ("Matched", "Zero-pay ($0.00)"):
            f = con.execute("SELECT * FROM mock_fee_schedule WHERE fee_row_key = ?", (r["matched_fee_row"],)).fetchone()
            assert f["mock_code"] == r["mock_code"] and f["claim_type"] == r["claim_type_requested"]
            assert f["effective_from"] <= r["service_date"] and (f["effective_to"] is None or r["service_date"] <= f["effective_to"])
        assert r["match_rows"] <= 1


def test_wrong_claim_type_and_out_of_range_are_distinct(con, loaded):
    tracker, _ = loaded
    for r in tracker:
        types = {x[0] for x in con.execute("SELECT claim_type FROM mock_fee_schedule WHERE mock_code = ?", (r["mock_code"],))}
        if r["bucket"] == "Wrong claim type":
            assert types and r["claim_type_requested"] not in types
        if r["bucket"] == "Out of date range":
            assert r["claim_type_requested"] in types
        if r["bucket"] == "No match":
            assert not types


def test_one_day_before_effective_date_is_out_of_range(loaded):
    tracker, _ = loaded
    lr052 = next(r for r in tracker if r["request_key"] == "LR052")
    assert (lr052["service_date"], lr052["earliest_effective_from"]) == ("2026-03-31", "2026-04-01")
    assert lr052["bucket"] == "Out of date range"


def test_duplicate_removed_once_and_original_kept(loaded):
    tracker, _ = loaded
    dups = [r for r in tracker if r["bucket"] == "Duplicate request"]
    assert len(dups) == 1
    original = next(r for r in tracker if r["request_key"] == dups[0]["duplicate_of"])
    assert original["bucket"] != "Duplicate request"


def test_python_classifier_agrees(con, loaded):
    tracker, _ = loaded
    assert M.classify_python(con) == {r["request_key"]: r["bucket"] for r in tracker}


def test_shortcut_errors_are_real_disagreements(loaded):
    tracker, naive = loaded
    errors = M.compare(tracker, naive)
    kinds = Counter(e["error"] for e in errors)
    assert kinds["Zero-pay read as not found"] == sum(1 for r in tracker if r["bucket"] == "Zero-pay ($0.00)")
    assert kinds["Priced under the wrong claim type"] > 0 and kinds["Priced from the wrong rate period"] > 0


def test_workbook_labels_every_amount_synthetic(loaded):
    from openpyxl import load_workbook
    tracker, _ = loaded
    wb = load_workbook(HERE / "fee_lookup_tracker.xlsx")
    assert wb.sheetnames == ["Tracker", "Summary", "Shortcut errors", "QA"]
    ws = wb["Tracker"]
    assert "SYNTHETIC" in ws["A1"].value and "SYNTHETIC" in ws["G3"].value
    for row in range(4, 4 + len(tracker)):
        assert ws.cell(row=row, column=7).number_format == M.AMOUNT_FORMAT
    assert wb["Summary"]["B4"].value.startswith("=COUNTIF(")


def test_qa_suite_has_no_failures(con, loaded):
    tracker, naive = loaded
    assert not [c for c in M.run_qa(con, tracker, naive, M.compare(tracker, naive)) if c["status"] == "fail"]
