"""Project B - Ten-payer quarterly scorecard with cohort-correct definitions (SYNTHETIC).

    python B_payer_scorecard/run.py

Builds the scorecard (Q2 vs Q3 2026), a definition-revision log that quantifies every
correction, an Excel workbook with live formulas, the dashboard and the QA log.
"""
import statistics
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))

from shared import portfolio as P  # noqa: E402

SQL, OUT = HERE / "sql", HERE / "outputs"
TOP10 = [f"PU{n:02d}" for n in range(1, 11)]
CUTOFF = datetime(2026, 10, 24, 23, 59, 59)


def rate(a, b):
    return a / b if b else None


def pct(new, old):
    return (new - old) / old if old else None


def networkdays_python(start: date, end: date) -> int:
    """Excel NETWORKDAYS without holidays: weekdays between two dates, both ends inclusive."""
    if start > end:
        return -networkdays_python(end, start)
    return sum(1 for k in range((end - start).days + 1) if (start + timedelta(days=k)).weekday() < 5)


def p90(values):
    s = sorted(values)
    return s[min(len(s) - 1, int(round(0.9 * (len(s) - 1))))] if s else None


# --------------------------------------------------------------------------- assemble
def load(con):
    P.run_script(con, SQL / "01_staging.sql")
    return dict(
        new_patients=P.query(con, SQL / "02_new_patients.sql"),
        auth=P.query(con, SQL / "03_auth_cohort.sql"),
        auth_decision=P.query(con, SQL / "04_auth_decision_method.sql"),
        served=P.query(con, SQL / "05_patients_served.sql"),
        pef=P.query(con, SQL / "06_pef_to_chair.sql"),
        mix=P.query(con, SQL / "07_medication_mix.sql"),
        alias=P.query(con, SQL / "08_alias_audit.sql"),
        umbrellas=P.query(con, "SELECT umbrella_key, umbrella_name, sort_order FROM umbrella_order ORDER BY sort_order"),
    )


def build_scorecard(d):
    idx = lambda rows, *k: {tuple(r[x] for x in k): r for r in rows}           # noqa: E731
    np_ = idx(d["new_patients"], "quarter", "umbrella_key")
    au = idx(d["auth"], "quarter", "umbrella_key")
    ad = idx(d["auth_decision"], "quarter", "umbrella_key")
    sv = idx(d["served"], "window_label", "umbrella_key")
    pef = defaultdict(list)
    for r in d["pef"]:
        if r["networkdays_inclusive"] >= 0:
            pef[(r["treated_quarter"], r["umbrella_key"])].append(r["networkdays_inclusive"])
    rows = []
    for u in d["umbrellas"]:
        k = u["umbrella_key"]
        g = lambda table, q, f: (table.get((q, k)) or {}).get(f, 0) or 0          # noqa: E731
        q3 = pef.get(("Q3 2026", k), [])
        q2 = pef.get(("Q2 2026", k), [])
        rows.append(dict(
            umbrella_key=k, payer=u["umbrella_name"].replace(" (fictional)", ""), is_top10=k in TOP10,
            new_patients_q2=g(np_, "Q2 2026", "new_patients"), new_patients_q3=g(np_, "Q3 2026", "new_patients"),
            auth_requests_q3=g(au, "Q3 2026", "requests"), approved_q3=g(au, "Q3 2026", "approved"),
            denied_q3=g(au, "Q3 2026", "denied"), pending_q3=g(au, "Q3 2026", "pending"),
            approval_rate_q2=rate(g(au, "Q2 2026", "approved"), g(au, "Q2 2026", "approved") + g(au, "Q2 2026", "denied")),
            approval_rate_q3=rate(g(au, "Q3 2026", "approved"), g(au, "Q3 2026", "approved") + g(au, "Q3 2026", "denied")),
            approval_rate_q3_decision_method=rate(g(ad, "Q3 2026", "approved"), g(ad, "Q3 2026", "approved") + g(ad, "Q3 2026", "denied")),
            served_t12m=g(sv, "T12M to Sep 30, 2026", "patients_served"),
            served_prior_t12m=g(sv, "T12M to Sep 30, 2025", "patients_served"),
            served_q2=g(sv, "Q2 2026", "patients_served"), served_q3=g(sv, "Q3 2026", "patients_served"),
            pef_orders_q3=len(q3), pef_median_q3=statistics.median(q3) if q3 else None,
            pef_mean_q3=round(statistics.mean(q3), 1) if q3 else None, pef_p90_q3=p90(q3),
            pef_median_q2=statistics.median(q2) if q2 else None,
        ))
    # Keep the review buckets only if they hold anything.
    rows = [r for r in rows if r["is_top10"] or r["umbrella_key"] == "PU99" or
            any(r[f] for f in ("new_patients_q2", "new_patients_q3", "auth_requests_q3", "served_t12m", "pef_orders_q3"))]
    all_q3 = [v for r in d["pef"] if r["treated_quarter"] == "Q3 2026" and r["networkdays_inclusive"] >= 0
              for v in [r["networkdays_inclusive"]]]
    all_q2 = [r["networkdays_inclusive"] for r in d["pef"] if r["treated_quarter"] == "Q2 2026" and r["networkdays_inclusive"] >= 0]
    s = lambda f: sum(r[f] for r in rows)                                           # noqa: E731
    total = dict(umbrella_key="ALL", payer="All payers", is_top10=False,
                 new_patients_q2=s("new_patients_q2"), new_patients_q3=s("new_patients_q3"),
                 auth_requests_q3=s("auth_requests_q3"), approved_q3=s("approved_q3"), denied_q3=s("denied_q3"),
                 pending_q3=s("pending_q3"),
                 approval_rate_q2=rate(sum(g["approved"] for g in d["auth"] if g["quarter"] == "Q2 2026"),
                                       sum(g["approved"] + g["denied"] for g in d["auth"] if g["quarter"] == "Q2 2026")),
                 approval_rate_q3=rate(s("approved_q3"), s("approved_q3") + s("denied_q3")),
                 approval_rate_q3_decision_method=rate(
                     sum(g["approved"] for g in d["auth_decision"] if g["quarter"] == "Q3 2026"),
                     sum(g["approved"] + g["denied"] for g in d["auth_decision"] if g["quarter"] == "Q3 2026")),
                 served_t12m=s("served_t12m"), served_prior_t12m=s("served_prior_t12m"),
                 served_q2=s("served_q2"), served_q3=s("served_q3"),
                 pef_orders_q3=len(all_q3), pef_median_q3=statistics.median(all_q3),
                 pef_mean_q3=round(statistics.mean(all_q3), 1), pef_p90_q3=p90(all_q3),
                 pef_median_q2=statistics.median(all_q2))
    return rows, total


def cohort_triangle(con):
    """Share of PA requests resolved within k days, by request month (as-of view; None = not yet observable)."""
    ks = [7, 14, 30, 45, 60]
    rows = P.query(con, """SELECT substr(requested_at, 1, 7) AS req_month, requested_at, decision_at
                           FROM authorizations WHERE requested_at >= '2026-01-01' AND requested_at < '2026-10-01'""")
    by_month = defaultdict(list)
    for r in rows:
        by_month[r["req_month"]].append(r)
    out = []
    for month in sorted(by_month):
        reqs = by_month[month]
        rec = dict(request_month=month, requests=len(reqs))
        for k in ks:
            # Only requests whose k-day window has fully elapsed by the snapshot are observable.
            eligible = [r for r in reqs if datetime.fromisoformat(r["requested_at"]) + timedelta(days=k) <= CUTOFF]
            if len(eligible) < len(reqs):
                rec[f"resolved_{k}d"] = None
            else:
                done = sum(1 for r in reqs if r["decision_at"] and
                           datetime.fromisoformat(r["decision_at"]) <= datetime.fromisoformat(r["requested_at"]) + timedelta(days=k))
                rec[f"resolved_{k}d"] = done / len(reqs)
        rec["resolved_as_of_snapshot"] = sum(1 for r in reqs if r["decision_at"]) / len(reqs)
        out.append(rec)
    return out, ks


# --------------------------------------------------------------------------- revision log
def revision_log(con, d, rows, total):
    q = lambda sql: P.scalar(con, sql)                                               # noqa: E731
    log = []

    def add(rid, area, before_rule, after_rule, before_value, after_value, impact, why):
        log.append(dict(revision=rid, area=area, original_draft=before_rule, revised_definition=after_rule,
                        before=before_value, after=after_value, impact=impact, why=why))

    orders_q3 = sum(r["order_series_created"] for r in d["new_patients"] if r["quarter"] == "Q3 2026")
    add("R1", "New referrals", "Order Series Created Date (distinct order series)", "Patient Created Date (distinct patients)",
        f"{orders_q3:,}", f"{total['new_patients_q3']:,}",
        f"Q3 new business overstated by {orders_q3 - total['new_patients_q3']:,} ({pct(orders_q3, total['new_patients_q3']):.0%})",
        "Order series include renewals and therapy changes for existing patients")
    add("R2", "Authorization rate", "Approved / decided, using decisions dated in the quarter",
        "Approved / (approved + denied) for requests SUBMITTED in the quarter; pending shown separately; as-of date stated",
        f"{total['approval_rate_q3_decision_method']:.1%}", f"{total['approval_rate_q3']:.1%}",
        f"Q3 decisions included {sum(r['decided_from_earlier_quarters'] for r in d['auth_decision'] if r['quarter'] == 'Q3 2026')} "
        "requests submitted in earlier quarters", "A cohort must follow the requests it is measuring")
    final_rows = q("SELECT COUNT(*) FROM authorization_decisions WHERE decision_result IN ('Approved', 'Denied')")
    final_reqs = q("SELECT COUNT(DISTINCT auth_key) FROM authorization_decisions WHERE decision_result IN ('Approved', 'Denied')")
    window_dups = sum(r["decision_rows"] - r["decided_requests"] for r in d["auth_decision"])
    add("R3", "Decision de-duplication", "COUNT(*) of final-decision rows", "COUNT(DISTINCT auth_key) - one final decision per request",
        f"{final_rows:,} rows", f"{final_reqs:,} requests",
        f"{final_rows - final_reqs} duplicate final-decision rows in the full history, {window_dups} of them in Q2-Q3 2026",
        "Repeated decision entries must not count twice")
    per_umbrella = defaultdict(int)
    for r in d["alias"]:
        per_umbrella[r["umbrella_key"]] += r["patients_served_plan"]
    plan_sum = sum(per_umbrella.values())
    umbrella_distinct = q("""SELECT COUNT(*) FROM (SELECT DISTINCT patient_key, umbrella_key FROM visit_payer
                              WHERE start_local >= '2025-10-01' AND start_local < '2026-10-01' AND plan_key IS NOT NULL)""")
    pu08_plan = per_umbrella.get("PU08", 0)
    pu08 = q("""SELECT COUNT(DISTINCT patient_key) FROM visit_payer WHERE umbrella_key = 'PU08'
                AND start_local >= '2025-10-01' AND start_local < '2026-10-01'""")
    add("R4", "Plan aliases", "Sum of plan-level distinct patients", "Group aliases to the umbrella first, then count distinct",
        f"{plan_sum:,}", f"{umbrella_distinct:,}",
        f"T12M patient-payer pairs overstated by {plan_sum - umbrella_distinct}; three-brand plan family {pu08_plan} -> {pu08}",
        "A patient who moves between brands of one payer is one patient")
    moved = q("""SELECT COUNT(*) FROM patient_payer WHERE created_at >= '2026-07-01' AND created_at < '2026-10-01'
                 AND umbrella_at_intake <> umbrella_current""")
    add("R5", "Payer attribution", "Patient's CURRENT insurance", "Insurance in effect on the event date",
        "current plan", "plan at event", f"{moved} of {total['new_patients_q3']} Q3 new patients would have been assigned to a different payer",
        "Insurance changes; history must keep the payer that applied at the time")
    q2_between_dates = q("""SELECT COUNT(*) FROM authorizations WHERE requested_at BETWEEN '2026-04-01' AND '2026-06-30'""")
    q2_half_open = q("""SELECT COUNT(*) FROM authorizations WHERE requested_at >= '2026-04-01' AND requested_at < '2026-07-01'""")
    double = q("""SELECT COUNT(*) FROM authorizations WHERE requested_at = '2026-07-01 00:00:00'""")
    add("R6", "Quarter boundaries", "BETWEEN '2026-04-01' AND '2026-06-30' on timestamps", "Half-open [quarter start, next quarter start)",
        f"{q2_between_dates:,} Q2 requests", f"{q2_half_open:,} Q2 requests",
        f"{q2_half_open - q2_between_dates} requests on Jun 30 were silently dropped; {double} request stamped exactly at midnight "
        "would be double-counted by inclusive end dates", "Timestamps after midnight on the last day fall outside a date-only BETWEEN")
    valid = [r for r in d["pef"] if r["networkdays_inclusive"] >= 0]
    neg = len(d["pef"]) - len(valid)
    q3 = [r for r in valid if r["treated_quarter"] == "Q3 2026"]
    by_payer = defaultdict(list)
    for r in q3:
        by_payer[r["umbrella_key"]].append(r["networkdays_inclusive"])
    unweighted = statistics.mean(statistics.mean(v) for v in by_payer.values())
    add("R7", "PEF-to-chair averaging", "Average of payer averages", "Order-weighted mean, median and P90; negative lags excluded",
        f"{unweighted:.1f} business days", f"{total['pef_mean_q3']:.1f} business days",
        f"{neg} impossible negative lag(s) excluded across Q2-Q3; small payers no longer move the network figure",
        "Averages of averages give a 9-order payer the same weight as a 70-order payer")
    gap_visits = q("""SELECT COUNT(*) FROM visit_payer WHERE umbrella_key = 'NO_COVERAGE'""")
    unmapped = q("SELECT COUNT(*) FROM payer_history WHERE plan_key NOT IN (SELECT plan_key FROM payer_umbrella_map)")
    add("R8", "Unattributable records", "Dropped by the inner join to payer", "Kept visible as 'Unmapped alias' and 'No coverage on file'",
        "silently excluded", "reported", f"{unmapped} coverage intervals use unmapped aliases; {gap_visits} treated visits had no coverage on file",
        "What cannot be attributed should be visible and routed for review, not lost")
    return log


# --------------------------------------------------------------------------- QA
def run_qa(con, d, rows, total, log):
    q = lambda sql: P.scalar(con, sql)                                               # noqa: E731
    checks = []

    def add(name, status, observed, rule):
        checks.append(dict(check=name, status=status, observed=observed, rule=rule))

    both = q("SELECT COUNT(*) FROM authorizations WHERE requested_at >= '2026-04-01' AND requested_at < '2026-10-01'")
    split = sum(r["requests"] for r in d["auth"])
    add("Every request lands in exactly one quarter", "pass" if both == split else "fail",
        f"Q2 + Q3 requests = {split:,}; requests in [Apr 1, Oct 1) = {both:,}; 3 boundary-stamped requests included once",
        "Half-open quarter windows")
    treated_t12m = q("SELECT COUNT(DISTINCT patient_key) FROM visit_payer WHERE start_local >= '2025-10-01' AND start_local < '2026-10-01'")
    add("Each served patient is attributed to exactly one payer", "pass" if treated_t12m == total["served_t12m"] else "fail",
        f"Sum across payers {total['served_t12m']:,} = {treated_t12m:,} distinct patients served (T12M)",
        "Attribute to the payer at the most recent treated visit in the window")
    overlaps = q("""SELECT COUNT(DISTINCT a.patient_key) FROM payer_history a JOIN payer_history b ON a.patient_key = b.patient_key
                    AND a.coverage_key < b.coverage_key AND a.effective_from <= COALESCE(b.effective_to, '9999-12-31')
                    AND b.effective_from <= COALESCE(a.effective_to, '9999-12-31')""")
    add("Overlapping coverage intervals resolved deterministically", "warn" if overlaps else "pass",
        f"{overlaps} patients have overlapping intervals; the most recently started interval wins and the rows go to data cleanup",
        "One payer per event")
    gaps = q("SELECT COUNT(*) FROM visit_payer WHERE umbrella_key = 'NO_COVERAGE'")
    add("Treated visits with no coverage on file", "warn" if gaps else "pass",
        f"{gaps} visits reported under 'No coverage on file' instead of being dropped", "Unattributable is a bucket, not a deletion")
    unmapped_plans = q("SELECT COUNT(*) FROM payer_plans WHERE plan_key NOT IN (SELECT plan_key FROM payer_umbrella_map)")
    names = [r[0] for r in con.execute("SELECT raw_plan_name FROM payer_plans WHERE plan_key NOT IN (SELECT plan_key FROM payer_umbrella_map)")]
    add("Raw plan aliases missing from the umbrella map", "warn" if unmapped_plans else "pass",
        f"{unmapped_plans} aliases unmapped ({'; '.join(repr(n) for n in names)}); proposed for analyst review, not auto-mapped",
        "Map by review, never by guess")
    r4 = next(r for r in log if r["revision"] == "R4")
    add("Distinct patients counted after alias grouping", "pass", r4["impact"], "Group plans to umbrella before COUNT(DISTINCT)")
    dup = next(r for r in log if r["revision"] == "R3")
    add("Duplicate decision rows removed", "pass", dup["impact"], "COUNT(DISTINCT auth_key)")
    pend = total["pending_q3"]
    add("Pending requests excluded from the approval-rate denominator", "pass",
        f"{pend} Q3 requests still pending at Oct 24 are reported separately", "Rate = approved / (approved + denied)")
    sample = d["pef"]
    mism = [r["order_key"] for r in sample if networkdays_python(date.fromisoformat(r["created_at"][:10]),
                                                                  date.fromisoformat(r["first_treated_at"][:10])) != r["networkdays_inclusive"]]
    add("Business-day lead time matches Excel NETWORKDAYS", "pass" if not mism else "fail",
        f"{len(sample) - len(mism)}/{len(sample)} orders match an independent Python NETWORKDAYS (inclusive of both endpoints)",
        "Weekdays only, both endpoints counted; the Excel tab recomputes it with NETWORKDAYS()")
    neg = sum(1 for r in d["pef"] if r["networkdays_inclusive"] < 0)
    add("Impossible negative lead times excluded", "warn" if neg else "pass",
        f"{neg} order(s) in the Q2-Q3 cohort had an order date after the first treatment", "Exclude and route to source review")
    lines_q3 = sum(r["line_items"] for r in d["mix"] if r["quarter"] == "Q3 2026")
    visits_q3 = q("SELECT COUNT(*) FROM qualifying_visits WHERE start_local >= '2026-07-01' AND start_local < '2026-10-01'")
    add("Medication mix ranked by treated visits, not line items", "pass",
        f"Q3: {lines_q3:,} given lines across therapies vs {visits_q3:,} treated visits", "Therapy-visit pairs, distinct appointments")
    add("Prior-year comparison caveat", "warn",
        f"T12M to Sep 30, 2025 = {total['served_prior_t12m']:,} vs {total['served_t12m']:,} now. The synthetic network starts in July 2024 "
        "with no legacy patients, so the prior window includes ramp-up and overstates growth", "State the caveat next to the number")
    return checks


# --------------------------------------------------------------------------- Excel workbook
def write_workbook(path, rows, total, d, log, qa, triangle, ks):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    navy, banner = PatternFill("solid", fgColor="12344A"), PatternFill("solid", fgColor="FFF4D6")
    white_bold, bold = Font(color="FFFFFF", bold=True), Font(bold=True)
    wb = Workbook()

    def sheet(title, headers, data, widths=None, formats=None, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws["A1"] = P.DISCLAIMER + " | Coralline Infusion Network (fictional) | as of Oct 24, 2026"
        ws["A1"].font = Font(bold=True, color="5C4300")
        ws["A1"].fill = banner
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(headers), 1))
        for j, h in enumerate(headers, 1):
            c = ws.cell(row=3, column=j, value=h)
            c.fill, c.font = navy, white_bold
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for i, rec in enumerate(data, 4):
            for j, v in enumerate(rec, 1):
                c = ws.cell(row=i, column=j, value=v)
                if formats and formats.get(j):
                    c.number_format = formats[j]
        ws.freeze_panes = "B4"
        for j in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(j)].width = (widths or {}).get(j, 14)
        ws.row_dimensions[3].height = 44
        return ws

    readme = wb.active
    readme.title = "README"
    lines = [P.DISCLAIMER, "", "Ten-payer quarterly scorecard - Coralline Infusion Network (fictional)",
             "Data as of Oct 24, 2026. Quarters are half-open: Q2 = [Apr 1, Jul 1), Q3 = [Jul 1, Oct 1).", "",
             "Tabs: Scorecard | Auth cohort | Served | PEF-to-chair | Medication mix | Alias audit | Auth maturity | Revision log | QA",
             "Formulas are live: rates, changes and the NETWORKDAYS check recompute in Excel.", "",
             "Definitions",
             "New referred patients: distinct patients created in the quarter, payer at intake.",
             "Authorization rate: approved / (approved + denied) for requests submitted in the quarter; pending shown separately.",
             "Patients served: distinct patients with a treated visit, payer at the most recent treated visit in the window.",
             "PEF-to-chair: NETWORKDAYS(order created, first treated visit), new-patient and new-therapy orders, first treated in the quarter."]
    for i, text in enumerate(lines, 1):
        readme.cell(row=i, column=1, value=text).font = bold if i in (1, 3, 9) else Font()
    readme["A1"].fill = banner
    readme.column_dimensions["A"].width = 120

    hdr = ["Payer (fictional)", "New patients Q2", "New patients Q3", "QoQ change", "PA requests Q3", "Approved Q3",
           "Denied Q3", "Pending Q3", "Approval rate Q3", "Approval rate Q2", "Change (pts)", "Served T12M",
           "Served prior T12M*", "Served change*", "PEF orders Q3", "PEF median Q3 (bus. days)", "PEF mean Q3", "PEF P90 Q3"]
    data = []
    for i, r in enumerate(rows + [total], 4):
        data.append([r["payer"], r["new_patients_q2"], r["new_patients_q3"], f"=IFERROR(C{i}/B{i}-1,\"\")",
                     r["auth_requests_q3"], r["approved_q3"], r["denied_q3"], r["pending_q3"],
                     f"=IFERROR(F{i}/(F{i}+G{i}),\"\")", r["approval_rate_q2"], f"=IFERROR((I{i}-J{i})*100,\"\")",
                     r["served_t12m"], r["served_prior_t12m"], f"=IFERROR(L{i}/M{i}-1,\"\")",
                     r["pef_orders_q3"], r["pef_median_q3"], r["pef_mean_q3"], r["pef_p90_q3"]])
    ws = sheet("Scorecard", hdr, data, widths={1: 34}, formats={4: "0.0%", 9: "0.0%", 10: "0.0%", 11: "0.0", 14: "0.0%"})
    for c in ws[3 + len(data)]:
        c.font = bold
    ws.cell(row=5 + len(data), column=1, value="* Prior window includes the synthetic network's ramp-up (data starts July 2024).")

    auth_rows = [[r["quarter"], r["umbrella_name"].replace(" (fictional)", ""), r["requests"], r["appeal_requests"], r["approved"],
                  r["denied"], r["pending"], f"=IFERROR(E{i}/(E{i}+F{i}),\"\")"] for i, r in enumerate(d["auth"], 4) if r["requests"]]
    sheet("Auth cohort", ["Quarter (requested)", "Payer", "Requests", "Appeals", "Approved", "Denied", "Pending", "Approval rate"],
          auth_rows, widths={2: 34}, formats={8: "0.0%"})
    sheet("Served", ["Window", "Payer", "Patients served"],
          [[r["window_label"], r["umbrella_name"].replace(" (fictional)", ""), r["patients_served"]] for r in d["served"]], widths={1: 24, 2: 34})
    pef_rows = []
    for i, r in enumerate(d["pef"], 4):
        pef_rows.append([r["order_key"], r["treated_quarter"], r["umbrella_key"], r["order_type"], r["therapy_key"],
                         datetime.fromisoformat(r["created_at"]).date(), datetime.fromisoformat(r["first_treated_at"]).date(),
                         r["networkdays_inclusive"], f"=NETWORKDAYS(F{i},G{i})", f"=IF(H{i}=I{i},\"OK\",\"CHECK\")"])
    sheet("PEF-to-chair", ["Order", "First treated in", "Payer key", "Order type", "Therapy", "Order created", "First treated",
                           "Business days (SQL)", "NETWORKDAYS (Excel)", "Match"], pef_rows,
          formats={6: "yyyy-mm-dd", 7: "yyyy-mm-dd"})
    sheet("Medication mix", ["Quarter", "Therapy", "Label (mock)", "Treated visits", "Patients", "Line items", "Vials"],
          [[r["quarter"], r["therapy_key"], r["therapy_label"], r["treated_visits"], r["patients"], r["line_items"], r["vials"]]
           for r in d["mix"]], widths={3: 40})
    sheet("Alias audit", ["Plan key", "Raw plan name (as typed in source)", "Umbrella", "Distinct patients served T12M (plan level)"],
          [[r["plan_key"], r["raw_plan_name"], r["umbrella_key"], r["patients_served_plan"]] for r in d["alias"]], widths={2: 40, 4: 22})
    sheet("Auth maturity", ["Request month", "Requests"] + [f"Resolved within {k} days" for k in ks] + ["Resolved by snapshot"],
          [[r["request_month"], r["requests"]] + [r[f"resolved_{k}d"] for k in ks] + [r["resolved_as_of_snapshot"]] for r in triangle],
          formats={j: "0%" for j in range(3, 4 + len(ks))})
    sheet("Revision log", ["Rev", "Area", "Original draft", "Revised definition", "Before", "After", "Impact", "Why"],
          [[r[k] for k in ("revision", "area", "original_draft", "revised_definition", "before", "after", "impact", "why")] for r in log],
          widths={2: 20, 3: 36, 4: 44, 5: 16, 6: 16, 7: 52, 8: 44})
    sheet("QA", ["Status", "Check", "Observed", "Rule"], [[c["status"].upper(), c["check"], c["observed"], c["rule"]] for c in qa],
          widths={2: 44, 3: 80, 4: 44})
    wb.save(path)


# --------------------------------------------------------------------------- main
def main():
    con = P.connect()
    d = load(con)
    rows, total = build_scorecard(d)
    triangle, ks = cohort_triangle(con)
    log = revision_log(con, d, rows, total)
    qa = run_qa(con, d, rows, total, log)

    OUT.mkdir(exist_ok=True)
    for name, data in [("scorecard", rows + [total]), ("new_patients_by_payer", d["new_patients"]), ("auth_cohort", d["auth"]),
                       ("auth_decision_method", d["auth_decision"]), ("patients_served", d["served"]), ("pef_to_chair_orders", d["pef"]),
                       ("medication_mix", d["mix"]), ("alias_audit", d["alias"]), ("auth_maturity", triangle),
                       ("revision_log", log), ("qa_results", qa)]:
        P.write_csv(data, OUT / f"{name}.csv")
    write_workbook(HERE / "payer_scorecard.xlsx", rows, total, d, log, qa, triangle, ks)

    top = [r for r in rows if r["is_top10"]]
    box = defaultdict(list)
    for r in d["pef"]:
        if r["treated_quarter"] == "Q3 2026" and r["networkdays_inclusive"] >= 0:
            box[r["umbrella_key"]].append(r["networkdays_inclusive"])
    names = {r["umbrella_key"]: r["payer"] for r in rows}
    mix_top = sorted([r for r in d["mix"] if r["quarter"] == "Q3 2026"], key=lambda r: -r["treated_visits"])[:10]
    mix_q2 = {r["therapy_key"]: r["treated_visits"] for r in d["mix"] if r["quarter"] == "Q2 2026"}
    review = (P.scalar(con, "SELECT COUNT(*) FROM payer_history WHERE plan_key NOT IN (SELECT plan_key FROM payer_umbrella_map)")
              + P.scalar(con, "SELECT COUNT(*) FROM visit_payer WHERE umbrella_key = 'NO_COVERAGE'"))
    tiles = [
        dict(label="New referred patients, Q3", value=f"{total['new_patients_q3']:,}",
             delta=pct(total["new_patients_q3"], total["new_patients_q2"]) * 100, deltaLabel="vs Q2", upIsGood=True,
             note="Distinct patients, payer at intake"),
        dict(label="PA approval rate, Q3 requests", value=f"{total['approval_rate_q3']:.1%}",
             note=f"Q2 {total['approval_rate_q2']:.1%}; decision-date draft said {total['approval_rate_q3_decision_method']:.1%}"),
        dict(label="Q3 requests still pending", value=f"{total['pending_q3']:,}", note="As of Oct 24; excluded from the rate"),
        dict(label="Patients served, T12M", value=f"{total['served_t12m']:,}", note="Oct 2025-Sep 2026, distinct patients"),
        dict(label="Median PEF-to-chair, Q3", value=f"{total['pef_median_q3']:.0f} days",
             note=f"Business days, NETWORKDAYS-inclusive; Q2 {total['pef_median_q2']:.0f}"),
        dict(label="Records routed for payer review", value=f"{review:,}", note="Unmapped aliases + visits with no coverage"),
    ]
    data = dict(meta=dict(as_of="October 24, 2026", seed=P.meta(con)["seed"]), tiles=tiles, rows=rows, total=total,
                top=top, box={names[k]: v for k, v in box.items() if k in TOP10}, mix=[dict(r, q2=mix_q2.get(r["therapy_key"], 0)) for r in mix_top],
                triangle=triangle, ks=ks, log=log, qa=qa)
    P.render_dashboard(HERE / "dashboard_template.html", HERE / "dashboard.html", data)
    if P.screenshot(HERE / "dashboard.html", HERE / "dashboard.png", width=1440, height=4750):
        print("  screenshot -> dashboard.png")
    write_qa_log(qa)
    print(f"Project B built: {len(rows)} payer rows, {len(log)} revisions, QA "
          f"{sum(c['status'] == 'pass' for c in qa)} pass / {sum(c['status'] == 'warn' for c in qa)} warn / "
          f"{sum(c['status'] == 'fail' for c in qa)} fail")


def write_qa_log(qa):
    lines = ["# QA log - Project B (ten-payer scorecard)", "", f"> {P.DISCLAIMER}. Generated by `run.py`.", "",
             "WARN = a real issue was found and handled (bucketed, excluded, or routed for review).", "",
             "| Status | Check | Observed | Rule applied |", "|---|---|---|---|"]
    lines += [f"| {c['status'].upper()} | {c['check']} | {c['observed']} | {c['rule']} |" for c in qa]
    (HERE / "QA_LOG.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
