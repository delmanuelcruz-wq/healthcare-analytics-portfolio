"""Project B tests: cohort definitions, attribution, alias grouping and lead-time math."""
import csv
import sys
from datetime import date
from pathlib import Path

import pytest

HERE = Path(__file__).resolve().parents[1]
ROOT = HERE.parent
sys.path.insert(0, str(ROOT))

from shared import portfolio as P   # noqa: E402
import importlib.util                 # noqa: E402
_spec = importlib.util.spec_from_file_location("project_b_run", HERE / "run.py")
B = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(B)


@pytest.fixture(scope="module")
def ctx():
    con = P.connect()
    d = B.load(con)
    rows, total = B.build_scorecard(d)
    log = B.revision_log(con, d, rows, total)
    yield dict(con=con, d=d, rows=rows, total=total, log=log)
    con.close()


@pytest.fixture(scope="module")
def register():
    with (ROOT / "foundation" / "data" / "exceptions_register.csv").open(encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def test_quarters_are_exclusive(ctx):
    con = ctx["con"]
    split = sum(r["requests"] for r in ctx["d"]["auth"])
    assert split == P.scalar(con, "SELECT COUNT(*) FROM authorizations WHERE requested_at >= '2026-04-01' AND requested_at < '2026-10-01'")


def test_boundary_requests_counted_once(ctx, register):
    con = ctx["con"]
    for auth in {r["record_key"] for r in register if r["defect_type"] == "QUARTER_BOUNDARY"}:
        ts = P.scalar(con, "SELECT requested_at FROM authorizations WHERE auth_key = ?", (auth,))
        hits = P.scalar(con, "SELECT COUNT(*) FROM quarters WHERE ? >= q_start AND ? < q_end", (ts, ts))
        assert hits == 1, (auth, ts)


def test_approval_rate_uses_resolved_requests_only(ctx):
    t = ctx["total"]
    assert t["approval_rate_q3"] == pytest.approx(t["approved_q3"] / (t["approved_q3"] + t["denied_q3"]))
    assert t["approved_q3"] + t["denied_q3"] + t["pending_q3"] == t["auth_requests_q3"]


def test_served_patients_attributed_once(ctx):
    con = ctx["con"]
    assert ctx["total"]["served_t12m"] == P.scalar(
        con, "SELECT COUNT(DISTINCT patient_key) FROM visit_payer WHERE start_local >= '2025-10-01' AND start_local < '2026-10-01'")


def test_alias_grouping_never_increases_counts(ctx):
    per_umbrella = {}
    for r in ctx["d"]["alias"]:
        per_umbrella.setdefault(r["umbrella_key"], 0)
        per_umbrella[r["umbrella_key"]] += r["patients_served_plan"]
    con = ctx["con"]
    for umbrella, plan_sum in per_umbrella.items():
        distinct = P.scalar(con, """SELECT COUNT(DISTINCT patient_key) FROM visit_payer WHERE umbrella_key = ?
                                    AND start_local >= '2025-10-01' AND start_local < '2026-10-01'""", (umbrella,))
        assert distinct <= plan_sum
    assert sum(per_umbrella.values()) > sum(
        P.scalar(con, """SELECT COUNT(DISTINCT patient_key) FROM visit_payer WHERE umbrella_key = ?
                         AND start_local >= '2025-10-01' AND start_local < '2026-10-01'""", (u,)) for u in per_umbrella)


def test_unmapped_and_uncovered_records_stay_visible(ctx, register):
    keys = {r["umbrella_key"] for r in ctx["rows"]}
    assert "NO_COVERAGE" in keys          # planted coverage gaps surface as their own bucket
    gap_visits = {r["related_key"] for r in register if r["defect_type"] == "COVERAGE_GAP"}
    found = {r[0] for r in ctx["con"].execute("SELECT appt_key FROM visit_payer WHERE umbrella_key = 'NO_COVERAGE'")}
    assert gap_visits == found


def test_networkdays_matches_excel_semantics(ctx):
    for r in ctx["d"]["pef"]:
        assert B.networkdays_python(date.fromisoformat(r["created_at"][:10]),
                                    date.fromisoformat(r["first_treated_at"][:10])) == r["networkdays_inclusive"]
    assert B.networkdays_python(date(2026, 10, 19), date(2026, 10, 23)) == 5     # Mon-Fri inclusive
    assert B.networkdays_python(date(2026, 10, 23), date(2026, 10, 26)) == 2     # Fri -> Mon
    assert B.networkdays_python(date(2026, 10, 21), date(2026, 10, 19)) == -3    # reversed, like Excel


def test_planted_negative_lead_times_are_excluded(ctx, register):
    planted = {r["record_key"] for r in register if r["defect_type"] == "NEGATIVE_LEAD_TIME"}
    negative = {r["order_key"] for r in ctx["d"]["pef"] if r["networkdays_inclusive"] < 0}
    assert planted == negative
    assert ctx["total"]["pef_orders_q3"] == sum(1 for r in ctx["d"]["pef"]
                                                if r["treated_quarter"] == "Q3 2026" and r["networkdays_inclusive"] >= 0)


def test_revision_log_quantifies_each_change(ctx):
    log = {r["revision"]: r for r in ctx["log"]}
    assert set(log) == {f"R{i}" for i in range(1, 9)}
    assert all(r["before"] != r["after"] for r in log.values())


def test_workbook_has_live_formulas():
    from openpyxl import load_workbook
    wb = load_workbook(HERE / "payer_scorecard.xlsx")
    assert {"Scorecard", "Revision log", "PEF-to-chair", "QA"} <= set(wb.sheetnames)
    assert str(wb["Scorecard"]["I4"].value).startswith("=IFERROR(")
    assert str(wb["PEF-to-chair"]["I4"].value).startswith("=NETWORKDAYS(")
    assert "SYNTHETIC" in wb["Scorecard"]["A1"].value
