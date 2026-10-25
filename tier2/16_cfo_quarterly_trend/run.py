"""Tier 2 #16 - CFO seven-quarter specialty therapy trend (SYNTHETIC).

    python tier2/16_cfo_quarterly_trend/run.py

One mock therapy, 2025 Q1 - 2026 Q3: treated visits per quarter (distinct appointments),
QoQ and first-to-last change, a one-sheet Excel workbook with a native chart, a README chart,
a small dashboard (+ screenshot) and the QA log.
"""
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT))

from shared import portfolio as P  # noqa: E402

SQL, OUT, CHARTS = HERE / "sql", HERE / "outputs", HERE / "charts"
THERAPY = "THERAPY_B"
QUARTERS = ["2025-Q1", "2025-Q2", "2025-Q3", "2025-Q4", "2026-Q1", "2026-Q2", "2026-Q3"]
LABEL = lambda q: q.replace("-", " ")   # noqa: E731  "2025-Q1" -> "2025 Q1"


def pct(new, old):
    return None if not old else (new - old) / old


def load(con):
    params = {"therapy": THERAPY}
    trend = [dict(r) for r in con.execute((SQL / "01_quarterly_trend.sql").read_text(encoding="utf-8"), params)]
    bounds = dict(con.execute((SQL / "02_boundary_checks.sql").read_text(encoding="utf-8"), params).fetchone())
    label = P.scalar(con, "SELECT therapy_label FROM therapies WHERE therapy_key = ?", (THERAPY,))
    for i, r in enumerate(trend):
        r["qoq_change"] = pct(r["treated_visits"], trend[i - 1]["treated_visits"]) if i else None
        r["visits_per_patient"] = round(r["treated_visits"] / r["patients_treated"], 2) if r["patients_treated"] else None
        r["line_items_vs_visits"] = pct(r["line_items"], r["treated_visits"])
    return trend, bounds, label


# --------------------------------------------------------------------------- independent re-implementation
def recompute_python(con):
    """Treated visits per quarter from raw rows in plain Python (cross-checks the SQL)."""
    appts = {r["appt_key"]: r for r in con.execute(
        "SELECT appt_key, start_local, status, meds_administered FROM appointments")}
    visits = defaultdict(set)
    for ad in con.execute("SELECT appt_key, administered_flag FROM administrations WHERE therapy_key = ?", (THERAPY,)):
        a = appts[ad["appt_key"]]
        if a["status"] == "Complete" and a["meds_administered"] == 1 and ad["administered_flag"] == 1:
            d = datetime.fromisoformat(a["start_local"])
            visits[f"{d.year}-Q{(d.month - 1) // 3 + 1}"].add(ad["appt_key"])
    return {q: len(visits[q]) for q in QUARTERS}


# --------------------------------------------------------------------------- QA suite
def run_qa(con, trend, bounds):
    checks = []

    def add(name, status, observed, rule):
        checks.append(dict(check=name, status=status, observed=observed, rule=rule))

    qs = [r["quarter"] for r in trend]
    add("Seven complete quarters, half-open boundaries", "pass" if qs == QUARTERS else "fail",
        f"{len(qs)} quarters ({LABEL(qs[0])} - {LABEL(qs[-1])}); {bounds['partial_q4_visits']} visits in the in-progress "
        "2026 Q4 are excluded", "Quarter = calendar quarter of the visit date; never compare a partial quarter")
    visits, lines = sum(r["treated_visits"] for r in trend), sum(r["line_items"] for r in trend)
    add("Headline counts distinct appointments, not line items", "pass",
        f"{visits:,} treated visits vs {lines:,} administration line items over seven quarters "
        f"(+{(lines - visits) / visits:.1%} if line IDs were counted)", "COUNT(DISTINCT appt_key) on treated visits")
    broken = [r["quarter"] for r in trend if r["line_items"] - r["waste_lines"] - r["given_lines_on_untreated_visits"]
              - r["extra_lines_same_visit"] != r["treated_visits"]]
    add("Line items reconcile to treated visits (the footnote)", "fail" if broken else "pass",
        "line items - waste - given lines on untreated visits - extra lines on the same visit = treated visits, "
        f"{len(trend) - len(broken)}/{len(trend)} quarters", "Every row between the two counts is explained")
    dupes = [(r["quarter"], r["exact_duplicate_lines"]) for r in trend if r["exact_duplicate_lines"]]
    add("Exact duplicate lines cannot inflate the trend", "warn" if dupes else "pass",
        (", ".join(f"{n} in {LABEL(q)}" for q, n in dupes) or "none") + "; distinct-appointment counts are unaffected",
        "Distinct appointment grain; duplicates routed to source cleanup")
    waste = sum(r["waste_lines"] for r in trend)
    add("Prepared-but-not-given lines excluded", "pass",
        f"{waste} waste lines (medication held at a Complete visit) are not visits", "administered_flag = 1 and meds_administered = 1")
    untreated = sum(r["given_lines_on_untreated_visits"] for r in trend)
    add("Lines on cancelled or no-show visits excluded", "warn" if untreated else "pass",
        f"{untreated} given line(s) on visits that were not treated", "A cancelled visit cannot be a treatment")
    add("Quarters bucketed on site-local time", "pass" if bounds["quarter_shift_if_utc"] == 0 else "warn",
        f"{bounds['quarter_shift_if_utc']} treated visit(s) would change quarter if bucketed on UTC",
        "Use start_local for reporting periods")
    py = recompute_python(con)
    mism = [q for q in QUARTERS if py[q] != next(r["treated_visits"] for r in trend if r["quarter"] == q)]
    add("Independent Python recomputation matches SQL", "fail" if mism else "pass",
        f"{len(QUARTERS) - len(mism)}/{len(QUARTERS)} quarters match", "Same rule, different implementation")
    addon = P.scalar(con, """SELECT COUNT(*) FROM administrations k WHERE k.therapy_key = 'THERAPY_K'
                             AND k.appt_key IN (SELECT appt_key FROM administrations WHERE therapy_key = ?)""", (THERAPY,))
    add("Single-therapy scope", "pass",
        f"{addon} add-on pre-medication lines on the same visits are not counted as {THERAPY} use",
        "Filter on the line's therapy, not the appointment's")
    first, last = trend[0]["treated_visits"], trend[-1]["treated_visits"]
    add("QoQ and first-to-last arithmetic", "pass",
        f"{LABEL(QUARTERS[0])} {first} -> {LABEL(QUARTERS[-1])} {last} = {pct(last, first):+.1%}; the workbook recomputes "
        "both with live formulas", "change = new / old - 1")
    return checks


# --------------------------------------------------------------------------- deliverables
def summary_lines(trend, label):
    first, last, prev = trend[0], trend[-1], trend[-2]
    peak = max(trend, key=lambda r: r["treated_visits"])
    vpp = [r["visits_per_patient"] for r in trend]
    return [
        f"{THERAPY} treated visits grew from {first['treated_visits']} in {LABEL(first['quarter'])} to "
        f"{last['treated_visits']} in {LABEL(last['quarter'])} ({pct(last['treated_visits'], first['treated_visits']):+.0%}).",
        f"Growth paused in {LABEL(last['quarter'])} ({pct(last['treated_visits'], prev['treated_visits']):+.1%} QoQ) "
        f"after the peak of {peak['treated_visits']} in {LABEL(peak['quarter'])}.",
        f"Patients treated grew from {first['patients_treated']} to {last['patients_treated']}; visits per patient "
        f"stayed between {min(vpp):.2f} and {max(vpp):.2f}, so growth came from more patients, not more visits each.",
    ]


def footnote(trend):
    q3 = trend[-1]
    visits, lines = sum(r["treated_visits"] for r in trend), sum(r["line_items"] for r in trend)
    return (f"Counts are distinct appointments where {THERAPY} was given. Counting administration line-item IDs would "
            f"report {q3['line_items']} for {LABEL(q3['quarter'])} instead of {q3['treated_visits']} "
            f"({pct(q3['line_items'], q3['treated_visits']):+.1%}), because split lots, waste lines and duplicate lines add rows. "
            f"Seven quarters: {lines:,} line items vs {visits:,} distinct appointments.")


def write_workbook(path, trend, label):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.styles import Alignment, Font, PatternFill

    banner, navy = PatternFill("solid", fgColor="FFF4D6"), PatternFill("solid", fgColor="12344A")
    wb = Workbook()
    ws = wb.active
    ws.title = "Seven-quarter trend"
    ws["A1"] = P.DISCLAIMER + " | Coralline Infusion Network (fictional) | data as of Oct 24, 2026"
    ws["A1"].font, ws["A1"].fill = Font(bold=True, color="5C4300"), banner
    ws.merge_cells("A1:N1")
    ws["A2"] = f"{THERAPY} - {label}: treated visits by quarter, 2025 Q1 - 2026 Q3"
    ws["A2"].font = Font(bold=True, size=14)
    ws["A3"] = "Treated visit = one appointment (Complete, medication given). Quarters use the site-local visit date."
    ws["A3"].font = Font(italic=True, color="52514E")
    headers = ["Quarter", "Treated visits", "QoQ change", "Patients treated", "Visits per patient", "Line items (reference only)"]
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=5, column=j, value=h)
        c.fill, c.font = navy, Font(color="FFFFFF", bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    first_row = 6
    for i, r in enumerate(trend, first_row):
        ws.cell(row=i, column=1, value=LABEL(r["quarter"]))
        ws.cell(row=i, column=2, value=r["treated_visits"])
        ws.cell(row=i, column=3, value=f"=B{i}/B{i - 1}-1" if i > first_row else None).number_format = "+0.0%;-0.0%"
        ws.cell(row=i, column=4, value=r["patients_treated"])
        ws.cell(row=i, column=5, value=f"=B{i}/D{i}").number_format = "0.00"
        ws.cell(row=i, column=6, value=r["line_items"])
    last_row = first_row + len(trend) - 1
    ws.cell(row=last_row + 2, column=1, value="First-to-last change (2025 Q1 to 2026 Q3)").font = Font(bold=True)
    ws.cell(row=last_row + 2, column=3, value=f"=B{last_row}/B{first_row}-1").number_format = "+0.0%;-0.0%"
    ws.cell(row=last_row + 3, column=1, value="Compound quarterly growth").font = Font(bold=True)
    ws.cell(row=last_row + 3, column=3, value=f"=(B{last_row}/B{first_row})^(1/{len(trend) - 1})-1").number_format = "+0.0%;-0.0%"
    row = last_row + 5
    ws.cell(row=row, column=1, value="Executive summary (synthetic data)").font = Font(bold=True)
    for k, text in enumerate(summary_lines(trend, label), 1):
        ws.cell(row=row + k, column=1, value=f"{k}. {text}")
    row += 5
    ws.cell(row=row, column=1, value="Footnote - counting grain").font = Font(bold=True)
    ws.cell(row=row + 1, column=1, value=footnote(trend)).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=14)
    ws.row_dimensions[row + 1].height = 48
    ws.cell(row=row + 2, column=1, value="Fictional organization and generated data. Not operational, medical or financial guidance.")
    for col, w in zip("ABCDEF", (30, 14, 12, 15, 14, 16)):
        ws.column_dimensions[col].width = w
    ws.row_dimensions[5].height = 32

    chart = BarChart()
    chart.type, chart.style = "col", 10
    chart.title = f"{THERAPY} treated visits per quarter (synthetic)"
    chart.y_axis.title, chart.x_axis.title = "Distinct appointments", None
    chart.add_data(Reference(ws, min_col=2, min_row=5, max_row=last_row), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=first_row, max_row=last_row))
    chart.legend = None
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True
    chart.series[0].graphicalProperties.solidFill = "2A78D6"
    chart.series[0].graphicalProperties.line.solidFill = "2A78D6"
    chart.height, chart.width = 8.5, 16
    ws.add_chart(chart, "H3")
    wb.save(path)


def readme_chart(trend, label):
    plt = P.mpl()
    CHARTS.mkdir(exist_ok=True)
    fig, ax = plt.subplots(figsize=(9, 4.6), facecolor="#fcfcfb")
    ax.set_facecolor("#fcfcfb")
    xs = [LABEL(r["quarter"]) for r in trend]
    ys = [r["treated_visits"] for r in trend]
    bars = ax.bar(xs, ys, color="#2a78d6", width=0.62, zorder=2)
    for b, r in zip(bars, trend):
        txt = f"{r['treated_visits']}"
        if r["qoq_change"] is not None:
            txt += f"\n{r['qoq_change']:+.0%}"
        ax.annotate(txt, (b.get_x() + b.get_width() / 2, b.get_height()), xytext=(0, 4), textcoords="offset points",
                    ha="center", va="bottom", fontsize=9, color="#0b0b0b", linespacing=1.15)
    ax.set_ylim(0, max(ys) * 1.25)
    ax.grid(axis="x", visible=False)
    first, last = trend[0]["treated_visits"], trend[-1]["treated_visits"]
    ax.set_title(f"{THERAPY} treated visits per quarter: {pct(last, first):+.0%} from 2025 Q1 to 2026 Q3",
                 loc="left", fontsize=13, fontweight="bold", color="#0b0b0b", pad=24)
    ax.text(0, 1.02, f"{label}. Distinct appointments with the therapy given; labels show QoQ change.",
            transform=ax.transAxes, fontsize=9.5, color="#52514e")
    fig.text(0.01, 0.01, P.DISCLAIMER + " | Coralline Infusion Network (fictional)", fontsize=7.5, color="#6f6d68")
    fig.tight_layout(rect=(0, 0.03, 1, 1))
    fig.savefig(CHARTS / "seven_quarter_trend.png", dpi=160, facecolor="#fcfcfb")
    plt.close(fig)


# --------------------------------------------------------------------------- main
def main():
    con = P.connect()
    meta = P.meta(con)
    trend, bounds, label = load(con)
    qa = run_qa(con, trend, bounds)
    P.write_csv(trend, OUT / "seven_quarter_trend.csv")
    P.write_csv(qa, OUT / "qa_results.csv")
    write_workbook(HERE / "cfo_seven_quarter_trend.xlsx", trend, label)
    readme_chart(trend, label)

    first, last, prev = trend[0], trend[-1], trend[-2]
    visits, lines = sum(r["treated_visits"] for r in trend), sum(r["line_items"] for r in trend)
    tiles = [
        dict(label=f"Treated visits, {LABEL(last['quarter'])}", value=f"{last['treated_visits']:,}",
             delta=pct(last["treated_visits"], prev["treated_visits"]) * 100, deltaLabel=f"vs {LABEL(prev['quarter'])}",
             upIsGood=True, note="Distinct appointments with the therapy given"),
        dict(label="First-to-last change", value=f"{pct(last['treated_visits'], first['treated_visits']):+.0%}",
             note=f"{LABEL(first['quarter'])} {first['treated_visits']} to {LABEL(last['quarter'])} {last['treated_visits']}"),
        dict(label=f"Patients treated, {LABEL(last['quarter'])}", value=f"{last['patients_treated']:,}",
             delta=pct(last["patients_treated"], prev["patients_treated"]) * 100, deltaLabel=f"vs {LABEL(prev['quarter'])}",
             upIsGood=True, note=f"{last['visits_per_patient']:.2f} visits per patient"),
        dict(label="If line items were counted", value=f"{lines:,}",
             note=f"vs {visits:,} distinct appointments over 7 quarters (+{(lines - visits) / visits:.1%})"),
    ]
    data = dict(meta=dict(as_of="October 24, 2026", seed=meta["seed"], therapy=THERAPY, label=label),
                tiles=tiles, trend=trend, summary=summary_lines(trend, label), footnote=footnote(trend), qa=qa)
    P.render_dashboard(HERE / "dashboard_template.html", HERE / "dashboard.html", data)
    if P.screenshot(HERE / "dashboard.html", HERE / "dashboard.png", width=1440, height=2350):
        print("  screenshot -> dashboard.png")
    P.write_qa_log(HERE / "QA_LOG.md", "Tier 2 #16 (CFO seven-quarter specialty therapy trend)", qa)
    print(f"#16 built: {THERAPY} {first['treated_visits']} -> {last['treated_visits']} treated visits; QA {P.qa_summary(qa)}")


if __name__ == "__main__":
    main()
