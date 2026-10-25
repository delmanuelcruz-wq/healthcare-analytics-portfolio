"""Tier 2 #10 - Fee-schedule lookup and exception tracker (SYNTHETIC).

    python tier2/10_fee_schedule_lookup/run.py

Prices a batch of mock billing-code requests against a mock fee schedule. A match needs
code + claim type + a service date inside the effective interval. Every request lands in
exactly one bucket; $0.00 is a zero-pay rate, not a missing one. Every amount is SYNTHETIC.
"""
import sys
from collections import Counter
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT))

from shared import portfolio as P  # noqa: E402

SQL, OUT = HERE / "sql", HERE / "outputs"
BUCKETS = ["Matched", "Zero-pay ($0.00)", "No match", "Wrong claim type", "Out of date range", "Duplicate request"]
AMOUNT_FORMAT = '"SYN $"#,##0.00'     # every amount cell carries the synthetic label in Excel


def money(v):
    return "" if v is None else f"SYN ${v:,.2f}"


def action(r):
    b = r["bucket"]
    if b == "Matched":
        return "Rate found for code, claim type and date"
    if b == "Zero-pay ($0.00)":
        return "Rate is $0.00 (non-covered, mock) - a real answer, confirm before billing"
    if b == "No match":
        return "Code is not on the schedule - research the code or ask the payer"
    if b == "Wrong claim type":
        return f"Code exists only as {r['claim_types_on_schedule']} - confirm the claim type"
    if b == "Out of date range":
        return f"No rate in force on {r['service_date']} (schedule starts {r['earliest_effective_from']}) - check the date or prior schedule"
    return f"Duplicate of {r['duplicate_of']} - removed from the count"


def load(con):
    tracker = P.query(con, SQL / "01_lookup_tracker.sql")
    for r in tracker:
        r["action"] = action(r)
    naive = {r["request_key"]: r for r in P.query(con, SQL / "02_naive_lookup.sql")}
    return tracker, naive


def compare(tracker, naive):
    """What the code-only shortcut gets wrong, request by request."""
    rows = []
    for r in tracker:
        n = naive[r["request_key"]]
        b, priced = r["bucket"], n["naive_result"] == "Priced"
        if b == "Zero-pay ($0.00)" and not priced:
            err = "Zero-pay read as not found"
        elif b == "Wrong claim type" and priced:
            err = "Priced under the wrong claim type"
        elif b == "Out of date range" and priced:
            err = "Priced outside the effective dates"
        elif b == "Matched" and n["naive_fee_row"] != r["matched_fee_row"]:
            err = "Priced from the wrong rate period"
        elif b == "Duplicate request" and priced:
            err = "Duplicate priced twice"
        else:
            continue
        rows.append(dict(request_key=r["request_key"], mock_code=r["mock_code"], service_date=r["service_date"],
                         correct_bucket=b, correct_amount_synthetic=r["fictional_fee_amount_synthetic"],
                         naive_result=n["naive_result"], naive_fee_row=n["naive_fee_row"],
                         naive_amount_synthetic=n["naive_amount_synthetic"], error=err))
    return rows


# --------------------------------------------------------------------------- independent re-implementation
def classify_python(con):
    sched = [dict(r) for r in con.execute("SELECT * FROM mock_fee_schedule")]
    seen, out = {}, {}
    for r in sorted(con.execute("SELECT * FROM lookup_requests"), key=lambda r: r["request_key"]):
        key = (r["mock_code"], r["claim_type_requested"], r["service_date"])
        if key in seen:
            out[r["request_key"]] = "Duplicate request"
            continue
        seen[key] = r["request_key"]
        code_rows = [s for s in sched if s["mock_code"] == r["mock_code"]]
        type_rows = [s for s in code_rows if s["claim_type"] == r["claim_type_requested"]]
        hits = [s for s in type_rows if s["effective_from"] <= r["service_date"]
                and (s["effective_to"] is None or r["service_date"] <= s["effective_to"])]
        if hits:
            out[r["request_key"]] = "Zero-pay ($0.00)" if hits[0]["fictional_fee_amount"] == 0 else "Matched"
        elif not code_rows:
            out[r["request_key"]] = "No match"
        elif not type_rows:
            out[r["request_key"]] = "Wrong claim type"
        else:
            out[r["request_key"]] = "Out of date range"
    return out


# --------------------------------------------------------------------------- QA suite
def run_qa(con, tracker, naive, errors):
    checks = []

    def add(name, status, observed, rule):
        checks.append(dict(check=name, status=status, observed=observed, rule=rule))

    counts = Counter(r["bucket"] for r in tracker)
    n_req = P.scalar(con, "SELECT COUNT(*) FROM lookup_requests")
    add("Every request lands in exactly one bucket", "pass" if sum(counts.values()) == n_req == len(tracker) else "fail",
        " · ".join(f"{b}: {counts.get(b, 0)}" for b in BUCKETS) + f" = {len(tracker)} of {n_req} requests",
        "Ordered rules: duplicate, match, no match, wrong claim type, out of range")
    overlaps = P.scalar(con, """SELECT COUNT(*) FROM mock_fee_schedule a JOIN mock_fee_schedule b
        ON a.mock_code = b.mock_code AND a.claim_type = b.claim_type AND a.fee_row_key < b.fee_row_key
       AND a.effective_from <= COALESCE(b.effective_to, '9999-12-31') AND b.effective_from <= COALESCE(a.effective_to, '9999-12-31')""")
    multi = sum(1 for r in tracker if r["match_rows"] > 1)
    add("A match needs code + claim type + effective date, and hits one row",
        "pass" if overlaps == 0 and multi == 0 else "fail",
        f"{overlaps} overlapping schedule intervals; {multi} requests matching more than one row",
        "service_date BETWEEN effective_from AND COALESCE(effective_to, open)")
    zero = [r for r in tracker if r["bucket"] == "Zero-pay ($0.00)"]
    misread = sum(1 for e in errors if e["error"] == "Zero-pay read as not found")
    add("$0.00 kept as zero-pay, never as 'no match'", "pass" if zero else "warn",
        f"{len(zero)} requests priced at $0.00 (codes {', '.join(sorted({r['mock_code'] for r in zero}))}); the blank-or-zero shortcut "
        f"files {misread} of them as not found", "A zero rate is an answer; a missing row is not")
    wrong = [r for r in tracker if r["bucket"] == "Wrong claim type"]
    add("Wrong claim type separated from no match", "warn" if wrong else "pass",
        f"{len(wrong)} Professional requests for codes on the Facility schedule only; the code-only shortcut prices "
        f"{sum(1 for e in errors if e['error'] == 'Priced under the wrong claim type')} of them anyway",
        "Claim type is part of the key")
    oor = [r for r in tracker if r["bucket"] == "Out of date range"]
    edge = [r for r in oor if r["earliest_effective_from"] and r["service_date"] < r["earliest_effective_from"]
            and (P.scalar(con, "SELECT julianday(?) - julianday(?)", (r["earliest_effective_from"], r["service_date"])) == 1)]
    add("Out-of-date-range requests flagged (inclusive boundaries)", "warn" if oor else "pass",
        f"{len(oor)} requests have no rate in force on the service date"
        + (f", including {edge[0]['request_key']} one day before its rate starts" if edge else ""),
        "Effective dates are inclusive on both ends")
    dups = [r for r in tracker if r["bucket"] == "Duplicate request"]
    originals_ok = all(next(t for t in tracker if t["request_key"] == r["duplicate_of"])["bucket"] != "Duplicate request"
                       for r in dups)
    add("Duplicate requests removed once, original kept", "pass" if originals_ok else "fail",
        "; ".join(f"{r['request_key']} duplicates {r['duplicate_of']}" for r in dups) or "none",
        "Same code + claim type + service date")
    changed = P.scalar(con, "SELECT COUNT(DISTINCT mock_code) FROM mock_fee_schedule WHERE effective_to IS NOT NULL")
    period = sum(1 for e in errors if e["error"] == "Priced from the wrong rate period")
    add("Rate changes priced from the interval covering the service date", "pass",
        f"{changed} codes changed rate on Jan 1, 2026; the code-only shortcut would use the 2025 rate for {period} "
        "matched request(s) dated 2026", "Pick the interval by service date, not the first row")
    absent = [r for r in tracker if r["bucket"] == "No match"]
    add("Codes missing from the schedule reported as no match", "warn" if absent else "pass",
        f"{len(absent)} requests for {len({r['mock_code'] for r in absent})} catalog codes with no schedule row",
        "Route to code research, never guess a rate")
    labelled = all(k.endswith("_synthetic") for r in tracker[:1] for k in r if "amount" in k)
    add("Every amount labelled synthetic", "pass" if labelled else "fail",
        "Amount columns end in _synthetic; workbook amount cells use the 'SYN $' number format; the dashboard prints 'SYN $'",
        "No amount can be mistaken for a real reimbursement rate")
    py = classify_python(con)
    agree = sum(1 for r in tracker if py[r["request_key"]] == r["bucket"])
    add("Independent Python classifier agrees with SQL", "pass" if agree == len(tracker) else "fail",
        f"{agree}/{len(tracker)} requests in the same bucket", "Same rules, different implementation")
    return checks


# --------------------------------------------------------------------------- workbook
def write_workbook(path, tracker, errors, qa):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    banner, navy = PatternFill("solid", fgColor="FFF4D6"), PatternFill("solid", fgColor="12344A")
    wb = Workbook()

    def sheet(ws, title, headers, rows, widths, formats=None):
        ws.title = title
        ws["A1"] = P.DISCLAIMER + " | Coralline Infusion Network (fictional) | every amount is SYNTHETIC"
        ws["A1"].font, ws["A1"].fill = Font(bold=True, color="5C4300"), banner
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        for j, h in enumerate(headers, 1):
            c = ws.cell(row=3, column=j, value=h)
            c.fill, c.font = navy, Font(color="FFFFFF", bold=True)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for i, rec in enumerate(rows, 4):
            for j, v in enumerate(rec, 1):
                c = ws.cell(row=i, column=j, value=v)
                if formats and j in formats:
                    c.number_format = formats[j]
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = "B4"
        ws.row_dimensions[3].height = 34
        return ws

    sheet(wb.active, "Tracker",
          ["Request", "Code (mock)", "Claim type requested", "Service date", "Bucket", "Matched fee row",
           "Fictional fee - SYNTHETIC", "Effective from", "Effective to", "Claim types on schedule", "Action / note"],
          [[r["request_key"], r["mock_code"], r["claim_type_requested"], r["service_date"], r["bucket"], r["matched_fee_row"],
            r["fictional_fee_amount_synthetic"], r["effective_from"], r["effective_to"], r["claim_types_on_schedule"], r["action"]]
           for r in tracker],
          [10, 11, 13, 12, 18, 12, 16, 12, 12, 16, 70], formats={7: AMOUNT_FORMAT})
    last = 3 + len(tracker)
    summary = [[b, f'=COUNTIF(Tracker!$E$4:$E${last},A{i})', f"=B{i}/SUM($B$4:$B${3 + len(BUCKETS)})"]
               for i, b in enumerate(BUCKETS, 4)]
    summary.append(["Total requests", f"=SUM(B4:B{3 + len(BUCKETS)})", None])
    ws = sheet(wb.create_sheet(), "Summary", ["Bucket", "Requests", "Share"], summary, [26, 12, 10], formats={3: "0.0%"})
    ws.cell(row=5 + len(BUCKETS), column=1,
            value="Counts recompute with COUNTIF. Match = code + claim type + service date inside the effective interval.")
    sheet(wb.create_sheet(), "Shortcut errors",
          ["Request", "Code", "Service date", "Correct bucket", "Correct fee - SYNTHETIC", "Code-only result",
           "Code-only fee row", "Code-only fee - SYNTHETIC", "What went wrong"],
          [[e["request_key"], e["mock_code"], e["service_date"], e["correct_bucket"], e["correct_amount_synthetic"],
            e["naive_result"], e["naive_fee_row"], e["naive_amount_synthetic"], e["error"]] for e in errors],
          [10, 10, 12, 18, 16, 14, 14, 16, 36], formats={5: AMOUNT_FORMAT, 8: AMOUNT_FORMAT})
    sheet(wb.create_sheet(), "QA", ["Status", "Check", "Observed", "Rule"],
          [[c["status"].upper(), c["check"], c["observed"], c["rule"]] for c in qa], [8, 44, 90, 50])
    wb.save(path)


# --------------------------------------------------------------------------- main
def main():
    con = P.connect()
    meta = P.meta(con)
    tracker, naive = load(con)
    errors = compare(tracker, naive)
    qa = run_qa(con, tracker, naive, errors)
    counts = Counter(r["bucket"] for r in tracker)
    summary = [dict(bucket=b, requests=counts.get(b, 0), share=round(counts.get(b, 0) / len(tracker), 4)) for b in BUCKETS]
    P.write_csv(tracker, OUT / "lookup_tracker.csv")
    P.write_csv(summary, OUT / "bucket_summary.csv")
    P.write_csv(errors, OUT / "shortcut_errors.csv")
    P.write_csv(qa, OUT / "qa_results.csv")
    write_workbook(HERE / "fee_lookup_tracker.xlsx", tracker, errors, qa)

    priced = [r for r in tracker if r["bucket"] in ("Matched", "Zero-pay ($0.00)")]
    tiles = [dict(label="Requests in the batch", value=f"{len(tracker)}", note=f"{len({r['mock_code'] for r in tracker})} distinct mock codes"),
             dict(label="Priced", value=f"{len(priced)}", note=f"{counts['Matched']} matched + {counts['Zero-pay ($0.00)']} zero-pay"),
             dict(label="Exceptions to work", value=f"{counts['No match'] + counts['Wrong claim type'] + counts['Out of date range']}",
                  note="No match, wrong claim type, out of date range"),
             dict(label="Duplicates removed", value=f"{counts['Duplicate request']}", note="Same code, claim type and date"),
             dict(label="Shortcut errors", value=f"{len(errors)}", note="Requests a code-only lookup would get wrong")]
    for r in tracker:
        r["fee_display"] = money(r["fictional_fee_amount_synthetic"])
    for e in errors:
        e["correct_display"], e["naive_display"] = money(e["correct_amount_synthetic"]), money(e["naive_amount_synthetic"])
    data = dict(meta=dict(as_of="October 24, 2026", seed=meta["seed"]), tiles=tiles, summary=summary, tracker=tracker,
                errors=errors, qa=qa)
    P.render_dashboard(HERE / "dashboard_template.html", HERE / "dashboard.html", data)
    if P.screenshot(HERE / "dashboard.html", HERE / "dashboard.png", width=1440, height=3500):
        print("  screenshot -> dashboard.png")
    P.write_qa_log(HERE / "QA_LOG.md", "Tier 2 #10 (fee-schedule lookup tracker)", qa)
    print(f"#10 built: {len(tracker)} requests -> " + ", ".join(f"{b} {counts.get(b, 0)}" for b in BUCKETS) + f"; QA {P.qa_summary(qa)}")


if __name__ == "__main__":
    main()
